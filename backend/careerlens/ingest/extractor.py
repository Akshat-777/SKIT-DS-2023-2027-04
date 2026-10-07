import pdfplumber
import fitz  # PyMuPDF
import docx
import time
from typing import Dict, Any, Tuple
from .ocr import extract_text_with_ocr

def extract_from_docx(file_path: str) -> Tuple[str, int]:
    """
    Extract text from DOCX files, including paragraphs, tables, headers, and footers.
    Returns extracted text and approximate page count (not perfectly accurate for DOCX).
    """
    doc = docx.Document(file_path)
    full_text = []
    
    # Headers and footers
    for section in doc.sections:
        if section.header:
            full_text.extend([p.text for p in section.header.paragraphs if p.text.strip()])
        if section.footer:
            full_text.extend([p.text for p in section.footer.paragraphs if p.text.strip()])
            
    # Main content and tables (simplified ordering, for resumes usually enough)
    for element in doc.element.body:
        if element.tag.endswith('p'):
            para = docx.text.paragraph.Paragraph(element, doc)
            if para.text.strip():
                full_text.append(para.text)
        elif element.tag.endswith('tbl'):
            table = docx.table.Table(element, doc)
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_data.append(cell.text.strip())
                if row_data:
                    full_text.append(" | ".join(row_data))
                    
    # Estimate pages by counting roughly 3000 chars per page
    text_result = "\n".join(full_text)
    pages = max(1, len(text_result) // 3000)
    return text_result, pages

def has_sufficient_text_layer(pdf_path: str) -> bool:
    """
    Determine if a PDF has a sufficient text layer by checking
    character count. If it's mostly images, return False.
    """
    try:
        with pdfplumber.open(pdf_path) as pdf:
            total_chars = 0
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    total_chars += len(text.strip())
            
            # If very few characters per page on average, it's likely scanned
            if len(pdf.pages) > 0 and total_chars / len(pdf.pages) < 100:
                return False
            return True
    except Exception:
        return False

def extract_from_pdf(file_path: str) -> Tuple[str, int, str]:
    """
    Extract text from PDF. Fallback to PyMuPDF if pdfplumber fails.
    """
    text = ""
    pages = 0
    method = "text_layer"
    try:
        with pdfplumber.open(file_path) as pdf:
            pages = len(pdf.pages)
            # Try extracting with reading order preservation where possible
            extracted_pages = []
            for page in pdf.pages:
                page_text = page.extract_text(x_tolerance=2, y_tolerance=3)
                if page_text:
                    extracted_pages.append(page_text)
            text = "\n\n".join(extracted_pages)
    except Exception:
        # Fallback to PyMuPDF
        doc = fitz.open(file_path)
        pages = len(doc)
        text = "\n\n".join([page.get_text() for page in doc])
        doc.close()
        
    return text, pages, method

def process_resume(file_path: str, original_filename: str, lang: str = 'eng') -> Dict[str, Any]:
    """
    Main extraction pipeline.
    """
    start_time = time.time()
    ext = original_filename.lower().split('.')[-1]
    
    text = ""
    pages = 0
    method = ""
    ocr_confidence = None
    warnings = []
    
    if ext == 'docx':
        text, pages = extract_from_docx(file_path)
        method = "text_layer"
    elif ext == 'pdf':
        if has_sufficient_text_layer(file_path):
            text, pages, method = extract_from_pdf(file_path)
        else:
            method = "ocr"
            text, ocr_confidence, pages = extract_text_with_ocr(file_path, lang)
            if ocr_confidence is not None and ocr_confidence < 60:
                warnings.append(f"Low OCR confidence ({ocr_confidence:.1f}%). Text might be inaccurate.")
    
    processing_ms = int((time.time() - start_time) * 1000)
    
    return {
        "text": text,
        "pages": pages,
        "method": method,
        "ocr_confidence": ocr_confidence,
        "warnings": warnings,
        "processing_ms": processing_ms
    }
